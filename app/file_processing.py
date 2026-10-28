import csv
import itertools
import os

from docx import Document
from pypdf import PdfReader


def read_file(path: str):
    ext = os.path.splitext(path)[1].lower()
    if os.path.getsize(path) == 0:
        filename = os.path.basename(path)
        raise ValueError(f"Empty file: {filename}")

    if ext == ".pdf":
        with open(path, "rb") as file:
            for i, page in enumerate(PdfReader(file).pages, 1):
                yield f"Page {i}", page.extract_text() or ""
    elif ext == ".docx":
        doc = Document(path)
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        yield "Document", "\n".join(lines)
    elif ext in (".txt", ".md"):
        with open(path, "r", encoding="utf-8", errors="replace") as file:
            yield "Document", file.read()
    elif ext == ".csv":
        with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as file:
            rows = csv.reader(file)
            headers = next(rows, None)
            if headers is None:
                return
            row_start = 1
            while batch := list(itertools.islice(rows, 500)):
                text = "\n".join(
                    [" | ".join(headers), *(" | ".join(row) for row in batch)]
                )
                row_end = row_start + len(batch) - 1
                yield f"Rows {row_start}-{row_end}", text
                row_start = row_end + 1
            if row_start == 1:
                yield "CSV", " | ".join(headers)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
